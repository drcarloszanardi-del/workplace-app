from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any


def recalc_pending_metrics_from_json(output_path: Path) -> bool:
    if '--recalc-pending-from-json' not in sys.argv:
        return False

    if not output_path.exists():
        raise SystemExit(
            f"Fallback requested but JSON not found: {output_path}. Restore src/data/flujo-fondos.json before using --recalc-pending-from-json."
        )

    payload = json.loads(output_path.read_text(encoding='utf8'))
    summary = payload.get('summary')
    summary_has_totals = isinstance(summary, list) and all(
        isinstance(item, dict) and isinstance(item.get('totals'), dict) for item in summary
    )
    records = payload.get('recordsPreview')

    if not summary_has_totals and not isinstance(records, list):
        raise SystemExit(
            'Fallback requested but src/data/flujo-fondos.json has neither summary totals nor a valid recordsPreview to rebuild pending metrics.'
        )

    metrics = payload.setdefault('metrics', {})
    if summary_has_totals:
        total_pending = round(
            sum(float(item['totals'].get('pending') or 0) for item in summary),
            2,
        )
        payload['pendingMetricsSource'] = 'summary_totals_fallback'
        # recordsPreview is capped, so it cannot safely rebuild openPendingCount.
        # Preserve the existing count unless a future fallback includes full records.
    else:
        open_pending = sum(1 for item in records if float(item.get('pending') or 0) > 0)
        total_pending = round(sum(float(item.get('pending') or 0) for item in records if float(item.get('pending') or 0) > 0), 2)
        metrics['openPendingCount'] = open_pending
        payload['pendingMetricsSource'] = 'records_preview_fallback'
    metrics['totalPending'] = total_pending
    payload['pendingMetricsRecalculatedFromJsonAt'] = datetime.now().isoformat()
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf8')
    print(f'Updated pending metrics in {output_path}')
    return True


def load_openpyxl_workbook():
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        missing = exc.name or 'openpyxl'
        raise SystemExit(
            f"Missing dependency: {missing}. Run `python3 -m venv /Users/jarvis/workplace-app/.venv && /Users/jarvis/workplace-app/.venv/bin/python -m pip install openpyxl && /Users/jarvis/workplace-app/.venv/bin/python scripts/build_flujo_data.py` outside the 5-minute cron, and do not commit `.venv/`. If that is not possible, update src/data/flujo-fondos.json manually only for metrics.totalPending as a temporary fallback. After that, validate with `python3 -m json.tool src/data/flujo-fondos.json >/dev/null` and inspect `git diff -- src/data/flujo-fondos.json scripts/build_flujo_data.py`."
        ) from exc
    return load_workbook

ROOT = Path('/Users/jarvis/workplace-app')
SOURCE = Path('/Users/jarvis/.openclaw/media/inbound/Flujo_de_Fondos_2026---5877ca3e-5fba-478c-bd25-d3518ab72ad8.xlsx')
OUTPUT = ROOT / 'src' / 'data' / 'flujo-fondos.json'
USAGE = f"""Usage:
  python3 scripts/build_flujo_data.py
  python3 scripts/build_flujo_data.py --recalc-pending-from-json

Notes:
- Full regeneration reads {SOURCE.name} and requires openpyxl.
- The fallback flag recalculates metrics.totalPending from the existing JSON, prefers summary totals when present, and leaves generatedAt untouched so the dashboard still shows that a full Excel regeneration is pending.
- The fallback only refreshes openPendingCount when summary totals are unavailable and the JSON must be rebuilt from recordsPreview.
- Do not run the full regeneration inside the 5-minute cron worker.
"""

if '--help' in sys.argv or '-h' in sys.argv:
    print(USAGE)
    raise SystemExit(0)

if recalc_pending_metrics_from_json(OUTPUT):
    raise SystemExit(0)

if not SOURCE.exists():
    raise SystemExit(
        f"Workbook not found: {SOURCE}. Restore the original Excel in that path before regenerating src/data/flujo-fondos.json, or update only metrics.totalPending manually as a temporary fallback."
    )

MONTH_NAMES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']


def as_number(value: Any) -> float:
    if value is None or value == '':
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        normalized = value.strip().replace('.', '').replace(',', '.')
        if not normalized:
            return 0.0
        try:
            return float(normalized)
        except ValueError:
            return 0.0
    return 0.0


def as_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    return None


load_workbook = load_openpyxl_workbook()
wb_values = load_workbook(SOURCE, data_only=True)
wb_formulas = load_workbook(SOURCE, data_only=False)

ws_datos_values = wb_values['Datos']
ws_resumen_values = wb_values['Flujo de Fondos']
ws_resumen_formulas = wb_formulas['Flujo de Fondos']
ws_deudas_values = wb_values['Deudas']

headers = [ws_datos_values.cell(1, c).value for c in range(1, 19)]
records = []
for r in range(2, ws_datos_values.max_row + 1):
    row = {headers[c - 1]: ws_datos_values.cell(r, c).value for c in range(1, 19)}
    if not any(value not in (None, '') for value in row.values()):
        continue
    record = {
        'id': r - 1,
        'date': as_date(row.get('Fecha1')),
        'month': int(as_number(row.get('Mes'))),
        'year': int(as_number(row.get('Año'))),
        'client': row.get('Cliente') or '',
        'product': row.get('Producto') or '',
        'category': row.get('RUBRO') or 'Sin rubro',
        'budget': round(as_number(row.get('Presupuesto')), 2),
        'supplierBudget': round(as_number(row.get('Presupuesto Proveedor')), 2),
        'collection1': round(as_number(row.get('Cobro 1')), 2),
        'pending': round(as_number(row.get('Pendiente')), 2),
        'date2': as_date(row.get('Fecha 2')),
        'month2': int(as_number(row.get('Mes2'))),
        'year2': int(as_number(row.get('Año2'))),
        'collection2': round(as_number(row.get('Cobro 2')), 2),
        'totalCollection': round(as_number(row.get('Cobro Total')), 2),
        'balance': round(as_number(row.get('Saldo')), 2),
        'expense': round(as_number(row.get('Gastos')), 2),
    }
    records.append(record)

month_columns = []
col = 2
while True:
    raw = ws_resumen_values.cell(2, col).value
    if not isinstance(raw, datetime):
        break
    month_columns.append({
        'col': col,
        'key': f'{raw.year}-{raw.month:02d}',
        'label': f"{MONTH_NAMES[raw.month - 1]} {str(raw.year)[2:]}",
        'year': raw.year,
        'month': raw.month,
    })
    col += 1


def category_map(start_row: int, end_row: int, month_col: int) -> dict[str, float]:
    values: dict[str, float] = {}
    for row in range(start_row, end_row + 1):
        label = ws_resumen_formulas.cell(row, 1).value
        if not label:
            continue
        values[str(label)] = round(as_number(ws_resumen_values.cell(row, month_col).value), 2)
    return values


summary = []
for month in month_columns:
    month_col = month['col']
    initial_balance = round(as_number(ws_resumen_values.cell(3, month_col).value), 2)
    income_by_category = category_map(5, 11, month_col)
    pending_by_category = category_map(13, 19, month_col)
    expense_by_category = category_map(21, 39, month_col)
    utility_by_category = category_map(41, 43, month_col)
    totals = {
        'income': round(as_number(ws_resumen_values.cell(4, month_col).value), 2),
        'pending': round(as_number(ws_resumen_values.cell(12, month_col).value), 2),
        'expense': round(as_number(ws_resumen_values.cell(20, month_col).value), 2),
        'utility': round(as_number(ws_resumen_values.cell(40, month_col).value), 2),
        'periodBalance': round(as_number(ws_resumen_values.cell(44, month_col).value), 2),
        'accumulatedBalance': round(as_number(ws_resumen_values.cell(45, month_col).value), 2),
    }
    summary.append({
        'key': month['key'],
        'label': month['label'],
        'year': month['year'],
        'month': month['month'],
        'initialBalance': initial_balance,
        'incomeByCategory': income_by_category,
        'pendingByCategory': pending_by_category,
        'expenseByCategory': expense_by_category,
        'utilityByCategory': utility_by_category,
        'totals': totals,
        'excelCheck': totals,
    })

sorted_records = sorted(records, key=lambda item: ((item['date'] or ''), item['id']), reverse=True)
open_pending = sum(1 for item in records if item['pending'] > 0)
total_pending = round(sum(item['pending'] for item in records if item['pending'] > 0), 2)
clients = len({item['client'] for item in records if item['client']})
categories = len({item['category'] for item in records if item['category']})

def top_items(key: str, limit: int = 8) -> list[dict[str, float | str]]:
    totals: dict[str, float] = {}
    for item in records:
        name = str(item.get(key) or '').strip()
        if not name:
            continue
        amount = item['totalCollection'] if key == 'client' else item['expense']
        if key == 'category' and amount == 0:
            amount = item['totalCollection']
        totals[name] = totals.get(name, 0.0) + float(amount)
    ranked = sorted(totals.items(), key=lambda pair: pair[1], reverse=True)[:limit]
    return [{'name': name, 'amount': round(amount, 2)} for name, amount in ranked]


debts = []
for row in range(2, 21):
    concept = ws_deudas_values.cell(row, 1).value
    amount = round(as_number(ws_deudas_values.cell(row, 2).value), 2)
    due_date = as_date(ws_deudas_values.cell(row, 3).value)
    if concept in (None, '') and amount == 0 and due_date is None:
        continue
    debts.append({
        'concept': concept or '',
        'amount': amount,
        'dueDate': due_date,
    })

output = {
    'sourceWorkbook': SOURCE.name,
    'generatedAt': datetime.now().isoformat(),
    'months': [{'key': m['key'], 'label': m['label'], 'year': m['year'], 'month': m['month']} for m in month_columns],
    'summary': summary,
    'metrics': {
        'recordCount': len(records),
        'clientsCount': clients,
        'categoriesCount': categories,
        'monthsCount': len(summary),
        'openPendingCount': open_pending,
        'totalPending': total_pending,
        'totalDebt': round(sum(item['amount'] for item in debts), 2),
    },
    'highlights': {
        'topClients': top_items('client'),
        'topIncomeCategories': top_items('category'),
        'topExpenseCategories': [
            {'name': name, 'amount': round(amount, 2)}
            for name, amount in sorted(
                ((cat, sum(item['expense'] for item in records if item['category'] == cat)) for cat in {item['category'] for item in records}),
                key=lambda pair: pair[1],
                reverse=True,
            )[:8]
        ],
    },
    'debts': debts,
    'recordsPreview': sorted_records[:150],
}

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding='utf8')
print(f'Wrote {OUTPUT}')
