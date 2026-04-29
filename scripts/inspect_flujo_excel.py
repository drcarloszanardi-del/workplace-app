from openpyxl import load_workbook
from collections import Counter
import json
from pathlib import Path

path = Path('/Users/jarvis/.openclaw/media/inbound/Flujo_de_Fondos_2026---5877ca3e-5fba-478c-bd25-d3518ab72ad8.xlsx')
wb = load_workbook(path, data_only=False)
summary = {}
for name in wb.sheetnames:
    ws = wb[name]
    data=[]
    for row in ws.iter_rows(values_only=True):
        vals=list(row)
        if any(v is not None and v != '' for v in vals):
            data.append(vals)
    summary[name] = {
        'rows': len(data),
        'cols': max((len(r) for r in data), default=0),
        'first_rows': data[:8],
    }

ws = wb['Datos']
headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
rows=[]
for r in range(2, ws.max_row+1):
    rec={headers[c-1]: ws.cell(r,c).value for c in range(1,ws.max_column+1)}
    rows.append(rec)

clientes=Counter()
productos=Counter()
for rec in rows:
    if rec.get('Cliente'):
        clientes[str(rec['Cliente']).strip()] += 1
    if rec.get('Rubro'):
        productos[str(rec['Rubro']).strip()] += 1
    elif len(headers) >= 17 and rec.get(headers[16]):
        productos[str(rec[headers[16]]).strip()] += 1

print(json.dumps({'summary':summary,'headers':headers,'top_clientes':clientes.most_common(20),'top_rubros':productos.most_common(30)}, ensure_ascii=False, indent=2, default=str))
